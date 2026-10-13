"""S14: statistics report, 10 slides (guide §9; plan §7.8, note §8).

Every number on a slide is read from outputs/stats (or the network summaries), never typed by hand. Data charts are
native PowerPoint charts; the co-authorship network drawing is the S12 figure. A PDF copy is made with LibreOffice
when it is installed.

Output: outputs/report/group29_report.pptx (+ .pdf)
"""
from __future__ import annotations

import json
import logging
import shutil
import subprocess

import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from . import StepFailed

log = logging.getLogger("dpcn.s14")

HEAD, BODY = "Cambria", "Calibri"
INK, INK2, MUTED, GRID = "1F2A2E", "4E5A5E", "8A9497", "DDE3E4"
DARK, PAPER, TINT = "16323A", "FFFFFF", "EEF4F4"
LAB_HEX = {"CSIR-CCMB": "2A78D6", "CSIR-IGIB": "EB6834", "CSIR-IMTECH": "1BAF7A", "CSIR-CDRI": "EDA100"}
Q_HEX = {"Q1": "104281", "Q2": "256ABF", "Q3": "5598E7", "Q4": "86B6EF"}
W, H = Inches(13.333), Inches(7.5)


ABBREV = [("National Institute of Pharmaceutical Education and Research", "NIPER"),
          ("All India Institute of Medical Sciences", "AIIMS"),
          ("Post Graduate Institute of Medical Education and Research", "PGIMER"),
          ("Indian Institute of Technology", "IIT"), ("Indian Institute of Chemical Technology", "CSIR-IICT"),
          ("Indraprastha Institute of Information Technology", "IIIT"), ("Biochemistry, Genetics and Molecular Biology", "Biochem., Genetics & Mol. Bio."),
          ("Agricultural and Biological Sciences", "Agri. & Biological Sciences"), ("Pharmacology, Toxicology and Pharmaceutics", "Pharmacology & Toxicology"),
          ("Immunology and Microbiology", "Immunology & Microbiology"), ("Institute", "Inst."), ("University", "Univ.")]


def abbr(name: str, n: int = 40) -> str:
    """Shorter display name (the CSVs keep the full one); cut at a word boundary with an ellipsis."""
    out = str(name)
    for long, short_ in ABBREV:
        if len(out) <= n:
            break
        out = out.replace(long, short_)
    return trunc(out, n)


def trunc(text: str, n: int) -> str:
    text = " ".join(str(text).split())
    return text if len(text) <= n else text[:n - 1].rsplit(" ", 1)[0] + "…"


def rgb(h): return RGBColor.from_string(h)
def short(lab): return lab.replace("CSIR-", "")
def pct(x, d=0): return "–" if x is None or pd.isna(x) else f"{100 * x:.{d}f}%"
def num(x): return "–" if x is None or pd.isna(x) else f"{int(round(x)):,}"


class Deck:
    def __init__(self):
        self.p = Presentation()
        self.p.slide_width, self.p.slide_height = W, H
        self.blank = self.p.slide_layouts[6]
        self.n = 0

    def slide(self, title: str, dark: bool = False, kicker: str | None = None):
        s = self.p.slides.add_slide(self.blank)
        self.n += 1
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(DARK if dark else PAPER)
        if not dark:
            self.text(s, title, 0.6, 0.35, 12.1, 0.8, size=30, bold=True, font=HEAD, color=INK)
            if kicker:
                self.text(s, kicker, 0.6, 1.05, 12.1, 0.45, size=15, color=INK2, italic=True)
            self.text(s, f"DPCN Group 29 · CSIR biology labs · {self.n}", 0.6, 7.05, 12.1, 0.3, size=10, color=MUTED,
                      align=PP_ALIGN.RIGHT)
        return s

    @staticmethod
    def text(s, text, x, y, w, h, size=14, bold=False, color=INK, font=BODY, italic=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.02)
        lines = text if isinstance(text, list) else [text]
        for i, line in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            r = para.add_run()
            r.text = line
            r.font.size, r.font.bold, r.font.italic, r.font.name = Pt(size), bold, italic, font
            r.font.color.rgb = rgb(color)
        return tb

    def bullets(self, s, items, x, y, w, h, size=14):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, it in enumerate(items):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            pPr = para._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.22)))
            pPr.set("indent", str(-Inches(0.22)))
            bu = pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar", {"char": "•"})
            pPr.append(bu)
            para.space_after = Pt(7)
            r = para.add_run()
            r.text = it
            r.font.size, r.font.name = Pt(size), BODY
            r.font.color.rgb = rgb(INK)
        return tb

    def card(self, s, x, y, w, h, fill=TINT):
        from pptx.enum.shapes import MSO_SHAPE
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.adjustments[0] = 0.06
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
        sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def stat(self, s, x, y, w, value, label, dark=False):
        self.card(s, x, y, w, 1.25, fill="21454F" if dark else TINT)
        self.text(s, value, x + 0.2, y + 0.12, w - 0.4, 0.6, size=28, bold=True, font=HEAD, color="FFFFFF" if dark else INK)
        self.text(s, label, x + 0.2, y + 0.72, w - 0.4, 0.45, size=12, color="C9D6D8" if dark else INK2)

    def table(self, s, df: pd.DataFrame, x, y, w, h, widths=None, size=11):
        rows, cols = len(df) + 1, len(df.columns)
        t = s.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h)).table
        for j, c in enumerate(df.columns):
            if widths:
                t.columns[j].width = Inches(widths[j])
        for i in range(rows):
            for j in range(cols):
                cell = t.cell(i, j)
                val = df.columns[j] if i == 0 else df.iat[i - 1, j]
                cell.text = "" if val is None or (isinstance(val, float) and pd.isna(val)) else str(val)
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(DARK if i == 0 else (TINT if i % 2 == 0 else PAPER))
                cell.margin_left = cell.margin_right = Inches(0.05)
                cell.margin_top = cell.margin_bottom = Inches(0.02)
                for para in cell.text_frame.paragraphs:
                    for r in para.runs:
                        r.font.size, r.font.name = Pt(size), BODY
                        r.font.bold = i == 0
                        r.font.color.rgb = rgb("FFFFFF" if i == 0 else INK)
        return t

    def chart(self, s, kind, categories, series: dict, x, y, w, h, colors=None, legend=True, number_format=None,
              labels=False, gap=60, stacked_100=False):
        cd = CategoryChartData()
        cd.categories = categories
        for name, vals in series.items():
            cd.add_series(name, vals)
        gf = s.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        ch = gf.chart
        ch.font.size, ch.font.name = Pt(11), BODY
        ch.font.color.rgb = rgb(INK2)
        ch.has_legend = legend and len(series) > 1
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
        try:
            va = ch.value_axis
            va.major_gridlines.format.line.color.rgb = rgb(GRID)
            va.format.line.fill.background()
            va.tick_labels.font.size = Pt(10)
            if number_format:
                va.tick_labels.number_format = number_format
                va.tick_labels.number_format_is_linked = False
            ca = ch.category_axis
            ca.format.line.color.rgb = rgb(GRID)
            ca.tick_labels.font.size = Pt(10)
        except (ValueError, AttributeError):
            pass
        plot = ch.plots[0]
        if hasattr(plot, "gap_width"):
            plot.gap_width = gap
        if stacked_100 and hasattr(plot, "overlap"):
            plot.overlap = 100
        for i, ser in enumerate(plot.series):
            c = (colors or [])[i] if colors and i < len(colors) else None
            if c:
                if kind in (XL_CHART_TYPE.LINE_MARKERS, XL_CHART_TYPE.LINE):
                    ser.format.line.color.rgb = rgb(c)
                    ser.format.line.width = Pt(2.25)
                    ser.marker.format.fill.solid()
                    ser.marker.format.fill.fore_color.rgb = rgb(c)
                    ser.marker.format.line.color.rgb = rgb(PAPER)
                    ser.smooth = False
                else:
                    ser.format.fill.solid()
                    ser.format.fill.fore_color.rgb = rgb(c)
        if labels:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.font.size = Pt(10)
            dl.font.color.rgb = rgb(INK)
            if number_format:
                dl.number_format = number_format
                dl.number_format_is_linked = False
            if kind in (XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.COLUMN_CLUSTERED):
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
        return ch


def load(cfg):
    st = cfg.paths.stats
    names = ["inst_summary", "inst_year_counts", "inst_quartile", "inst_top_journals", "faculty_metrics", "hindex_hist",
             "inst_fields", "top_cited_papers", "dept_metrics", "citation_pagerank_top", "citation_flow_year",
             "citation_flow_summary", "top_citing_institutes", "authorship_centrality", "communities", "hub_authors",
             "partner_ranking", "cross_cluster_summary", "lab_citation_matrix", "core_ranks"]
    missing = [n for n in names if not (st / f"{n}.csv").exists()]
    if missing:
        raise StepFailed(f"missing stats tables {missing} - run s09 to s12 first")
    T = {n: pd.read_csv(st / f"{n}.csv") for n in names}
    T["indeg"] = pd.read_csv(st / "citation_indegree_dist.csv")
    T["cit"] = json.loads((st / "citation_network_summary.json").read_text())
    T["au"] = json.loads((st / "authorship_network_summary.json").read_text())
    T["meta"] = json.loads((cfg.paths.submission / "metadata.json").read_text())
    T["s05"] = json.loads((cfg.paths.interim / "s05_summary.json").read_text())
    T["s06"] = json.loads((cfg.paths.interim / "s06_summary.json").read_text())
    return T


def build(cfg, T) -> Presentation:
    labs = [l["short"] for l in cfg.labs]
    S = T["inst_summary"].set_index("scope")
    cagr = next(c for c in S.columns if c.startswith("cagr_"))
    d = Deck()
    years = list(range(cfg.window.start, cfg.window.end + 1))

    # 1 ---- title
    s = d.slide("", dark=True)
    d.text(s, "DPCN Research Mapping · Group 29", 0.7, 0.7, 12, 0.5, size=16, color="9FC3C9")
    d.text(s, "Research output, citation and collaboration networks of four CSIR biology labs, 2016–2026", 0.7, 1.2,
           11.9, 1.7, size=34, bold=True, font=HEAD, color="FFFFFF")
    d.text(s, "CSIR-CCMB Hyderabad · CSIR-IGIB New Delhi · CSIR-IMTECH Chandigarh · CSIR-CDRI Lucknow", 0.7, 3.25, 12, 0.5,
           size=16, color="C9D6D8")
    a = S.loc["ALL"]
    for i, (v, l) in enumerate([(num(a.papers), "unique papers"), (num(a.citations), "citations"),
                                (num(a.faculty_with_papers), "faculty with papers"), (pct(a.q1_fraction), "Q1 share of ranked journal papers"),
                                (num(a.institutional_h), "institutional h-index")]):
        d.stat(s, 0.7 + i * 2.45, 4.3, 2.25, v, l, dark=True)
    d.text(s, f"Source: OpenAlex snapshot {T['meta']['collection_date']}; Scimago SJR; CORE2023. Instructor: Chittaranjan Hens. "
              f"{cfg.window.partial_year} is a partial year.", 0.7, 6.6, 12, 0.4, size=11, color="9FC3C9")

    # 2 ---- data and method
    s = d.slide("Data and method", kicker="A reproducible OpenAlex pipeline; every number here is regenerated from the submitted CSVs")
    steps = [("Roster", f"{num(S.loc['ALL'].roster_faculty)} scientists: lab websites plus PIs found from lab-affiliated authorship"),
             ("Match", "Each scientist matched to one or more OpenAlex author IDs (ORCID, own publication list, name)"),
             ("Filter", "A paper counts only if the scientist's own authorship carries the lab; articles, reviews, letters, conference papers"),
             ("Enrich", f"SJR quartile of the publication year (coverage {pct(min(v['coverage'] for v in T['s06']['labs'].values()))}–{pct(max(v['coverage'] for v in T['s06']['labs'].values()))}); OpenAlex topics; Indian citation flows")]
    for i, (h_, t_) in enumerate(steps):
        x = 0.6 + i * 3.08
        d.card(s, x, 1.75, 2.88, 2.25)
        d.text(s, f"{i + 1}  {h_}", x + 0.2, 1.88, 2.5, 0.45, size=18, bold=True, font=HEAD)
        d.text(s, t_, x + 0.2, 2.38, 2.5, 1.55, size=13, color=INK2)
    cov = pd.DataFrame(T["s05"]["coverage"])
    cov["Lab"] = cov["lab"].map(short)
    cov = cov.merge(S[["faculty_with_papers", "papers"]].reset_index(), left_on="lab", right_on="scope")
    d.table(s, pd.DataFrame({"Lab": cov["Lab"], "Faculty": cov["faculty_with_papers"].map(num),
                             "Papers": cov["papers"].map(num),
                             "Coverage of lab output": cov["coverage_%"].map(lambda v: f"{v:.0f}%")}),
            0.6, 4.35, 6.4, 1.9, widths=[1.4, 1.3, 1.4, 2.3], size=12)
    d.bullets(s, ["Faculty = scientists/PIs; students and postdocs appear only as co-authors",
                  "h-index is computed within 2016–2026 from this dataset (lifetime h shown for context)",
                  "Indian partner = an Indian institution other than the paper's own lab (CSIR HQ, ministries, AcSIR excluded)",
                  "Limitations: OpenAlex author disambiguation; uncovered papers are mostly by students or unlisted staff"],
              7.4, 4.3, 5.4, 2.6, size=13)

    # 3 ---- trend
    yc = T["inst_year_counts"]
    full = yc[(yc.scope.isin(labs)) & (yc.year < cfg.window.partial_year)]
    peaks = full.loc[full.groupby("scope")["papers"].idxmax()].set_index("scope")["year"]
    biggest = S.loc[labs, "papers"].idxmax()
    s = d.slide("Publication trend 2016–2026",
                kicker=f"{short(biggest)} publishes most · peak years: " +
                       ", ".join(f"{short(l)} {peaks[l]}" for l in labs) + f" · {cfg.window.partial_year} is a partial year")
    d.chart(s, XL_CHART_TYPE.LINE_MARKERS, [str(y) + ("*" if y == cfg.window.partial_year else "") for y in years],
            {short(l): yc[yc.scope == l].sort_values("year")["papers"].tolist() for l in labs}, 0.5, 1.6, 8.4, 5.3,
            colors=[LAB_HEX[l] for l in labs])
    for i, l in enumerate(labs):
        r = S.loc[l]
        d.card(s, 9.2, 1.7 + i * 1.3, 3.55, 1.15)
        d.text(s, short(l), 9.4, 1.78 + i * 1.3, 1.6, 0.4, size=15, bold=True, font=HEAD, color=LAB_HEX[l])
        d.text(s, f"{num(r.papers)} papers · {num(r.citations)} citations", 9.4, 2.15 + i * 1.3, 3.3, 0.35, size=12, color=INK2)
        d.text(s, f"CAGR {cfg.window.start}–{cfg.window.cagr_end}: {pct(r[cagr], 1)}", 9.4, 2.45 + i * 1.3, 3.3, 0.35, size=12, color=INK2)

    # 4 ---- journal quality
    s = d.slide("Journal quality: Q1 fraction and top journals",
                kicker="Q1 fraction = Q1 papers ÷ journal papers with a known SJR quartile")
    q = T["inst_quartile"]
    order = (labs + ["ALL"])[::-1]          # bar charts draw the first category at the bottom
    cats = [short(l) if l != "ALL" else "All" for l in order]
    d.chart(s, XL_CHART_TYPE.BAR_STACKED_100, cats,
            {qq: [float(q[(q.scope == sc) & (q.quartile == qq)]["share_of_known"].iloc[0]) for sc in order]
             for qq in ["Q1", "Q2", "Q3", "Q4"]}, 0.5, 1.6, 5.6, 4.2, colors=[Q_HEX[k] for k in Q_HEX],
            number_format="0%", gap=50, stacked_100=True)
    d.text(s, "  ".join(f"{short(l)} Q1 {pct(S.loc[l].q1_fraction)} (coverage {pct(S.loc[l].quartile_coverage)})" for l in labs),
           0.6, 5.95, 5.6, 0.8, size=12, color=INK2)
    tj = T["inst_top_journals"]
    tj = tj[tj.scope != "ALL"]
    rows = []
    for l in labs:
        g = tj[tj.scope == l].head(3)
        for r in g.itertuples():
            rows.append({"Lab": short(l) if r.rank == 1 else "", "Journal": trunc(r.journal_name.title() if r.journal_name.isupper() else r.journal_name, 48),
                         "Papers": r.papers, "Q": r.quartile})
    d.text(s, "Top 3 journals per lab (top 5 in the dashboard)", 6.5, 1.55, 6.3, 0.4, size=14, bold=True, font=HEAD)
    d.table(s, pd.DataFrame(rows), 6.5, 2.0, 6.3, 4.6, widths=[0.9, 4.1, 0.75, 0.55], size=11)

    # 5 ---- faculty
    s = d.slide("Faculty: h-index distribution and leaders",
                kicker="Window h-index from 2016–2026 dataset papers; department-level summary from the roster divisions")
    hh = T["hindex_hist"]
    hh = hh[hh["bin"] != "summary"]
    bins = sorted(hh["bin_start"].unique())
    bins = [b for b in bins if hh[hh.bin_start >= b]["faculty"].sum() > 0]
    d.chart(s, XL_CHART_TYPE.COLUMN_STACKED, [f"{b}–{b + 4}" for b in bins],
            {short(l): [int(hh[(hh.scope == l) & (hh.bin_start == b)]["faculty"].sum()) for b in bins] for l in labs},
            0.5, 1.6, 6.2, 4.4, colors=[LAB_HEX[l] for l in labs], gap=30)
    summ = T["hindex_hist"][T["hindex_hist"]["bin"] == "summary"].set_index("scope")
    d.text(s, "Median h: " + " · ".join(f"{short(l)} {summ.loc[l, 'median']:.0f}" for l in labs) +
           f" · all {summ.loc['ALL', 'median']:.0f}", 0.6, 6.1, 6.1, 0.4, size=12, color=INK2)
    fm = T["faculty_metrics"].sort_values(["h_window", "citations"], ascending=False).head(10)
    d.table(s, pd.DataFrame({"Faculty": fm["name"], "Lab": fm["lab"].map(short), "Papers": fm["papers"],
                             "Citations": fm["citations"].map(num), "h": fm["h_window"]}),
            7.0, 1.6, 5.8, 3.6, widths=[2.4, 0.9, 0.8, 1.0, 0.7], size=11)
    dm = T["dept_metrics"].sort_values("papers", ascending=False)
    top_dept = dm.iloc[0]
    d.bullets(s, [f"{len(dm)} departments/divisions; largest by output: {top_dept.department} ({short(top_dept.lab)}, {num(top_dept.papers)} papers)",
                  "IMTECH publishes no divisions: grouped by each scientist's main OpenAlex subfield"],
              7.0, 5.4, 5.8, 1.5, size=12)

    # 6 ---- research areas & impact
    s = d.slide("Research areas and most-cited papers",
                kicker="OpenAlex primary topic → field; citation counts at collection date")
    f = T["inst_fields"]
    top_fields = [x for x in f[f.scope == "ALL"].sort_values("papers", ascending=False)["field"] if x != "Other"][:6]
    d.chart(s, XL_CHART_TYPE.BAR_CLUSTERED, [abbr(x, 30) for x in top_fields][::-1],
            {short(l): [float(f[(f.scope == l) & (f.field == x)]["share"].sum()) for x in top_fields][::-1] for l in labs},
            0.5, 1.6, 5.4, 5.3, colors=[LAB_HEX[l] for l in labs], number_format="0%", gap=40)
    tc = T["top_cited_papers"]
    tc = tc[tc.scope == "ALL"]
    d.table(s, pd.DataFrame({"#": tc["rank"], "Paper": tc["title"].map(lambda t: trunc(t, 72)), "Journal": tc["journal"].map(lambda t: trunc(t, 26)),
                             "Year": tc["year"], "Cites": tc["citations"].map(num), "Lab": tc["labs"].str.replace("CSIR-", "")}),
            6.2, 1.6, 6.6, 5.3, widths=[0.35, 3.15, 1.35, 0.55, 0.65, 0.55], size=9)

    # 7 ---- citation network
    c = T["cit"]
    s = d.slide("Citation network within the four labs",
                kicker="Directed graph: nodes = dataset papers, edges = paper A cites paper B (internal corpus)")
    for i, (v, l) in enumerate([(num(c["nodes"]), "papers (nodes)"), (num(c["internal_edges"]), "internal citations"),
                                (pct(c["largest_wcc"] / c["nodes"]), "in largest component"),
                                (pct(c["share_indeg_0"]), "never cited internally")]):
        d.stat(s, 0.6 + i * 1.6, 1.6, 1.45, v, l)
    m = T["lab_citation_matrix"].set_index(T["lab_citation_matrix"].columns[0])
    mm = m.copy()
    mm.index = [short(x) for x in mm.index]
    mm.columns = [short(x) for x in mm.columns]
    mm = mm.reset_index().rename(columns={"index": "cites →"})
    d.table(s, mm.astype(str), 0.6, 3.2, 6.2, 1.8, widths=[1.2, 1.25, 1.25, 1.25, 1.25], size=12)
    deg = sorted(np.repeat(T["indeg"]["in_degree"].values, T["indeg"]["papers"].values), reverse=True)
    top10 = sum(deg[:max(1, len(deg) // 10)]) / max(sum(deg), 1)
    mv = m.values
    same_lab = mv.trace() / mv.sum() if mv.sum() else 0
    d.bullets(s, [f"Mean in-degree {c['mean_indeg']:.2f}; maximum {c['max_indeg']}; reciprocity {c['reciprocity']:.3f}",
                  f"Heavy-tailed: the most-cited 10% of papers receive {pct(top10)} of internal citations",
                  f"{pct(same_lab)} of lab-to-lab citations stay inside the citing lab (matrix diagonal)"],
              0.6, 5.25, 6.2, 1.7, size=12)
    pr = T["citation_pagerank_top"]
    pr = pr[pr.scope == "ALL"].head(10)
    d.text(s, "PageRank top 10", 7.2, 1.6, 5.6, 0.4, size=15, bold=True, font=HEAD)
    d.table(s, pd.DataFrame({"Paper": pr["title"].fillna("").map(lambda t: trunc(t, 64)), "Year": pr["year"],
                             "Lab": pr["labs"].str.replace("CSIR-", ""), "In-deg": pr["indeg"]}),
            7.2, 2.05, 5.6, 4.8, widths=[3.6, 0.6, 0.75, 0.65], size=9)

    # 8 ---- Indian citation flow
    s = d.slide("Indian citation flow", kicker="Citations received from (in) and given to (out) works carrying another Indian institution")
    fy = T["citation_flow_year"]
    d.chart(s, XL_CHART_TYPE.LINE_MARKERS, [str(y) + ("*" if y == cfg.window.partial_year else "") for y in years],
            {f"{short(l)}": [int(fy[(fy.lab == l) & (fy.year == y)]["in"].sum()) for y in years] for l in labs},
            0.5, 1.6, 6.4, 3.9, colors=[LAB_HEX[l] for l in labs])
    d.text(s, f"Citations received per citing year ({cfg.window.partial_year}* partial)", 0.6, 5.5, 6.2, 0.35, size=12, color=INK2, italic=True)
    fs = T["citation_flow_summary"].set_index("lab")
    for i, l in enumerate(labs):
        d.stat(s, 0.6 + i * 1.6, 5.85, 1.45, f"{fs.loc[l, 'impact_ratio']:.2f}", f"{short(l)} in/out ratio")
    tci = T["top_citing_institutes"]
    rows = []
    for l in labs:
        for j, r in enumerate(tci[tci.your_inst == l].head(3).itertuples()):
            rows.append({"Lab": short(l) if j == 0 else "", "Citing institute": abbr(r.partner_inst, 44), "Cites in": num(r.cites_in)})
    d.text(s, "Top Indian institutes citing each lab", 7.3, 1.6, 5.5, 0.4, size=15, bold=True, font=HEAD)
    d.table(s, pd.DataFrame(rows), 7.3, 2.05, 5.5, 4.0, widths=[0.9, 3.6, 1.0], size=10)
    d.text(s, "In = " + ", ".join(f"{short(l)} {num(fs.loc[l, 'cites_in'])}" for l in labs) + "; out = " +
           ", ".join(f"{short(l)} {num(fs.loc[l, 'cites_out'])}" for l in labs), 7.3, 6.25, 5.5, 0.6, size=11, color=INK2)

    # 9 ---- authorship network
    au = T["au"]["faculty_graph"]
    s = d.slide("Authorship network: brokers and communities",
                kicker=f"Faculty co-authorship graph: {au['nodes']} faculty, {au['edges']} links, Louvain modularity Q = {au['modularity']:.2f}")
    img = cfg.paths.figures / "authorship_faculty.png"
    if img.exists():
        s.shapes.add_picture(str(img), Inches(0.5), Inches(1.55), height=Inches(5.4))
    ce = T["authorship_centrality"]
    top_deg = ce.sort_values("degree", ascending=False).head(5)
    top_btw = ce.sort_values("betweenness", ascending=False).head(5)
    d.text(s, "Most collaborative (degree)", 6.9, 1.6, 2.9, 0.4, size=14, bold=True, font=HEAD)
    d.table(s, pd.DataFrame({"Faculty": top_deg["name"], "Lab": top_deg["lab"].map(short), "Deg": top_deg["degree"]}),
            6.9, 2.0, 2.9, 2.0, widths=[1.75, 0.65, 0.5], size=10)
    d.text(s, "Bridges (betweenness)", 9.95, 1.6, 2.9, 0.4, size=14, bold=True, font=HEAD)
    d.table(s, pd.DataFrame({"Faculty": top_btw["name"], "Lab": top_btw["lab"].map(short),
                             "Betw": top_btw["betweenness"].map(lambda v: f"{v:.3f}")}),
            9.95, 2.0, 2.9, 2.0, widths=[1.7, 0.65, 0.55], size=10)
    hub = T["hub_authors"].iloc[0]
    sizes = au["community_sizes"]
    cm = T["communities"]
    single = cm[cm["size"] >= 3]["labs"].map(lambda x: x.count("(") == 1).mean()
    d.bullets(s, [f"{au['communities']} communities; the largest have {sizes[0]}, {sizes[1]} and {sizes[2]} faculty; "
                  f"{pct(single)} of communities with 3+ members sit within a single lab",
                  f"{au['cross_lab_edges']} of {au['edges']} faculty links join two labs; {au['isolates']} faculty have no faculty co-author",
                  f"Hub author: {hub['name']} ({short(hub['lab'])}) co-authors with {int(hub['n_partner_insts'])} Indian partner institutions",
                  f"All-author graph: {T['au']['internal_graph']['nodes']:,} authors, Q = {T['au']['internal_graph']['modularity']:.2f}; "
                  f"{T['au']['papers_skipped_over_max_authors']} papers with > {T['au']['max_authors_for_pairs']} authors excluded from pairs"],
              6.9, 4.3, 5.95, 2.7, size=12)

    # 10 ---- connectivity map & takeaways
    s = d.slide("Connectivity map and key findings", kicker="Who our labs co-author with across Indian science")
    cc = T["cross_cluster_summary"]
    agg = cc.groupby("pair_label")["shared_papers"].sum().sort_values(ascending=False)
    agg = agg[agg > 0].head(8)[::-1]
    d.chart(s, XL_CHART_TYPE.BAR_CLUSTERED, [x.replace("↔", " – ") for x in agg.index], {"Shared papers": [int(v) for v in agg.values]},
            0.5, 1.6, 5.6, 4.0, colors=["16323A"], legend=False, labels=True, gap=40)
    d.text(s, "Shared papers by partner cluster (all four labs)", 0.6, 5.6, 5.5, 0.35, size=12, color=INK2, italic=True)
    pr_ = T["partner_ranking"]
    rows = []
    for l in labs:
        r = pr_[pr_.your_inst == l].iloc[0]
        rows.append({"Lab": short(l), "Top partner": abbr(r.partner_inst, 36), "Papers": int(r.shared_papers),
                     "Jaccard": "–" if pd.isna(r.jaccard_topics) else f"{r.jaccard_topics:.2f}"})
    d.table(s, pd.DataFrame(rows), 0.6, 6.0, 5.5, 0.95, widths=[0.8, 3.1, 0.75, 0.85], size=10)
    best = S.loc[labs].sort_values("q1_fraction", ascending=False)
    biggest = S.loc[labs].sort_values("papers", ascending=False)
    ext = cc[~cc.pair_label.str.endswith(("↔CSIR", "↔Other"))].groupby("pair_label")["shared_papers"].sum().sort_values(ascending=False)
    gives = [short(l) for l in labs if fs.loc[l, "impact_ratio"] < 1]
    gets = [short(l) for l in labs if fs.loc[l, "impact_ratio"] >= 1]
    flow_line = (f"Indian citation in/out ratio {fs['impact_ratio'].min():.2f}–{fs['impact_ratio'].max():.2f}: "
                 + (f"{', '.join(gets)} receive more than they give" if gets else "")
                 + ("; " if gets and gives else "") + (f"{', '.join(gives)} give more than they receive" if gives else "") + ".")
    findings = [
        f"{short(biggest.index[0])} is the most productive lab ({num(biggest.iloc[0].papers)} papers); "
        f"{short(best.index[0])} has the highest Q1 share ({pct(best.iloc[0].q1_fraction)}).",
        f"{pct(S.loc['ALL'].pct_indian_partner_collab)} of papers have an Indian partner co-author and "
        f"{pct(S.loc['ALL'].pct_foreign_collab)} a foreign one.",
        "Strongest ties outside CSIR: " + ", ".join(f"{k.split('↔')[1]} ({num(v)} papers)" for k, v in ext.head(3).items()) + ".",
        flow_line,
        f"Collaboration is lab-bound: {au['cross_lab_edges']} of {au['edges']} faculty links and {pct(1 - same_lab)} of "
        "lab-to-lab citations cross between the four labs.",
    ]
    d.card(s, 6.5, 1.6, 6.3, 5.35)
    d.text(s, "Key findings", 6.75, 1.75, 5.8, 0.45, size=18, bold=True, font=HEAD)
    d.bullets(s, findings, 6.75, 2.3, 5.85, 4.5, size=15)
    return d.p


def run(cfg, oa, argv=()) -> None:
    T = load(cfg)
    prs = build(cfg, T)
    if len(prs.slides) > 10:
        raise StepFailed(f"report has {len(prs.slides)} slides; the limit is 10")
    out_dir = cfg.paths.report
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"group{cfg.group}_report.pptx"
    prs.save(path)
    print(f"\n{cfg.rel(path)}: {len(prs.slides)} slides")
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(path)],
                       capture_output=True, timeout=300)
        if (out_dir / f"group{cfg.group}_report.pdf").exists():
            print(f"{cfg.rel(out_dir / f'group{cfg.group}_report.pdf')} written")
    else:
        log.warning("LibreOffice not found; export the PDF from PowerPoint by hand")
